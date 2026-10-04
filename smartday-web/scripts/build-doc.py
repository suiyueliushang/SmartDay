# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = r"C:\Users\liushang\OneDrive\Major\project\calendar-dsh"
SHOTS = os.path.join(ROOT, "smartday-web", ".docs-shots")
DSHOTS = os.path.join(ROOT, "smartday-desktop", ".smoke")
OUT = os.path.join(ROOT, "SmartDay功能介绍文档.docx")
ACCENT = RGBColor(0x4F, 0x6E, 0xF7)
GREY = RGBColor(0x66, 0x66, 0x66)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "微软雅黑"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.8)

def h(text, level=1, size=None):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.color.rgb = ACCENT
    run.font.size = Pt(size or {1: 17, 2: 13.5, 3: 11.5}.get(level, 11))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

def para(text, size=10.5, color=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(6)

def bullets(items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def figure(path, caption, width=6.4):
    if not os.path.exists(path):
        para("[缺少截图: " + os.path.basename(path) + "]", color=GREY, italic=True)
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    cap.paragraph_format.space_after = Pt(12)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SmartDay · 智能日程与任务管理")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = ACCENT
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("功能介绍文档")
r.font.size = Pt(15)
r.font.color.rgb = GREY
doc.add_paragraph()
para("一个把「日历、任务、笔记、专注、提醒」装进同一处的个人效率工具：", size=11)
bullets([
    "电脑桌面就是日历 —— 壁纸级桌面日历，不动手就能看见今天与本周要做什么；"
    "一天的事都在一页里 —— 日程、待办、笔记、倒计时同屏查看；"
    "提醒一定会到你手上 —— 浏览器通知之外，还能同时推送到 QQ 机器人、邮箱、微信；"
    "数据存在你自己的设备上 —— 无需注册、没有服务端，网页版与桌面端共用同一份数据。"
])
para("本文档面向第一次接触 SmartDay 的读者，包含功能说明与真实运行界面截图。", size=10, color=GREY, italic=True)
doc.add_page_break()

h("一、功能总览")
para("SmartDay 由 8 个模块组成，覆盖「计划 → 执行 → 记录 → 复盘」的完整闭环：")
table(["模块", "一句话说明", "核心能力"], [
    ["🖥️ 桌面壁纸日历", "把日历变成桌面壁纸", "透明置底、点击穿透、锁定/编辑双模式、4 主题 2 材质、托盘常驻"],
    ["🏠 概览", "今天一眼看清", "迷你日历、今日概览、当天汇总、纪念日倒计时、本周统计"],
    ["📅 日历", "月/周/日/年/议程 五视图", "农历节日、班休角标、分类筛选、重复事件、单击看当天、双击新建"],
    ["✅ 任务", "分组 × 清单 × 拖拽", "分组/清单管理、优先级、子任务、我的一天、四象限、看板、日历视图"],
    ["📝 笔记", "日记与笔记合并", "标签筛选、时间线、Markdown 阅读/编辑、自动保存"],
    ["🎯 专注", "番茄钟 + 数据复盘", "番茄钟/正向计时、任务绑定、热力图、年/月/周/日分析、历史回填"],
    ["🔔 提醒与通知", "多通道同时送达", "事件/任务/纪念日提醒、免打扰、QQ 机器人 + 邮箱 + 微信、推送记录"],
    ["⚙️ 设置与数据", "可控、可迁移", "主题明暗、备份导出导入、回收站、全局搜索、云同步骨架"],
])
doc.add_page_break()

h("二、桌面壁纸日历（桌面端特色）")
para("这是 SmartDay 最有辨识度的功能：日历直接铺在桌面上，像壁纸一样常驻，但它是活的 ——"
     "今天做了哪些事、下一个安排是什么、离纪念日还有几天，不用打开任何软件就能看到。")
bullets([
    "置底显示：贴在壁纸层，不挡窗口、不抢焦点、不进任务栏；"
    "点击穿透：空白处点一下直接落到桌面图标，只有按钮与事件才响应鼠标；"
    "锁定 / 编辑双模式：🔒 锁定后不可误改；解锁即可拖动、缩放、调透明度；"
    "4 套主题 × 2 种材质：墨、雾、砂、薄荷，玻璃或纯色，透明度 0–100 连续可调；"
    "托盘常驻：双击托盘图标打开主应用；锁定时窗口立即回到最底层。"
])
figure(os.path.join(DSHOTS, "wallpaper.png"), "图 1  桌面壁纸日历（真实桌面效果）")
figure(os.path.join(DSHOTS, "wallpaper-menu.png"), "图 2  ☰ 菜单：主题、材质、透明度、缩放、锁定")
figure(os.path.join(DSHOTS, "theme-ink.png"), "图 3  另一套主题（墨）")
doc.add_page_break()

h("三、主应用功能详解")
h("3.1 概览：今天一眼看清", 2)
para("打开主应用的第一屏。左侧是迷你日历与今日安排，右侧是纪念日倒计时与本周统计，底部是「当天汇总」——点日历上的任意一天，内容会跟着切换。")
bullets([
    "迷你日历：有安排的日期带圆点，周末、节假日着色；"
    "今日概览：按时间列出今日日程与任务，勾选即可完成；"
    "当天汇总：那一天的事件、任务、日记数量一览，点击可跳转；"
    "纪念日与倒计时：正数倒数到纪念日，负数显示已经过去多少天。"
])
figure(os.path.join(SHOTS, "01-overview.png"), "图 4  概览页：迷你日历 + 今日概览 + 当天汇总 + 纪念日 + 本周统计")
h("3.2 日历：五种视图，农历节日班休俱全", 2)
bullets([
    "五视图：月 / 周 / 日 / 年 / 议程，一键切换；"
    "日期信息：阳历、农历（如「廿四」）、节日（国庆节）、节气（霜降）、调休「班」与「休」角标；"
    "周末着色：周六周日单独配色，与桌面日历保持一致；"
    "分类筛选：分类卡可勾选显示/隐藏，右键可设为默认；"
    "单击日期：侧栏显示当天的事件、任务与笔记，点击条目分别打开事件弹窗、任务详情、笔记；"
    "双击日期：直接新建事件（默认 9:00，时长可在设置中调整）；"
    "重复事件、拖拽改期、周数显示（W40/W41）一应俱全。"
])
figure(os.path.join(SHOTS, "02-calendar.png"), "图 5  月视图：单击日期后显示「当天详情」，选中日期带蓝色描边")
h("3.3 任务：清单、分组、四象限与拖拽归类", 2)
bullets([
    "组织方式：分组 → 清单 → 任务 三层，支持新建/重命名/删除，清单与分组都能拖拽排序；"
    "拖拽归类：把任务直接拖到左侧分组或清单上即可完成归类；"
    "任务详情：点击任务从右侧滑出详情面板，可改标题、日期、优先级、子任务、备注；"
    "智能视图：今天、明天、最近七天、我的一天、已计划、已完成、全部；"
    "多视角：列表 / 看板 / 四象限（重要紧急）/ 日历；"
    "重复任务：每天、每周、每月、每年，完成后自动滚动到下一个周期。"
])
figure(os.path.join(SHOTS, "03-tasks.png"), "图 6  任务页：左侧分组与清单，中间任务列表，右侧详情面板")
h("3.4 笔记：日记与笔记合并为一种", 2)
bullets([
    "一个页面搞定：左侧「全部笔记 / 全部标签」，右侧笔记流；"
    "时间信息：每条笔记显示「创建时间」与「最后修改时间」，修改不会重置创建时间；"
    "阅读优先：点击列表先进入阅读视图，再点「编辑」才进入编辑，避免误改；"
    "Markdown 支持：标题、列表、引用、加粗、代码块，编辑与预览分栏；"
    "自动保存：输入即存，不会丢内容，也不会重复生成笔记；"
    "标签体系：按标签一键筛选，标签数量实时显示；"
    "日记：按日期记录心情，从日历/当天详情可直达当天日记。"
])
figure(os.path.join(SHOTS, "04-notes.png"), "图 7  笔记页：左侧「全部笔记 / 全部标签」，右侧笔记流（含创建与最后修改时间）")
h("3.5 专注：番茄钟 + 可复盘的数据", 2)
bullets([
    "两种计时：番茄钟（默认 25 分钟，可自定义）与正向计时，时长就在开始卡片里设置；"
    "必须点「开始专注」：不会一进页面就自动计时；"
    "绑定任务：专注可关联任务或事件，专注时长会累计到该任务；"
    "同时只允许一个专注：其它入口会提示「已有专注进行中」；"
    "到点提醒：桌面通知 + 提示音，并写入通知中心（因此也会推送到手机）；"
    "专注分析：热力图（类似 GitHub 贡献图）、年度/月度/周/日四种口径、趋势柱状图、按任务/按清单排行；"
    "历史记录：可按状态与来源筛选、关键字搜索、导出 CSV；已完成的记录还能重新绑定任务。"
])
figure(os.path.join(SHOTS, "05-focus.png"), "图 8  专注页：开始专注卡片（含时长与关联任务）+ 专注分析热力图")
figure(os.path.join(SHOTS, "focus-analysis-month.png"), "图 9  按月度分析：该月的专注热力图与统计")
h("3.6 提醒与通知：一定能送到你手上", 2)
para("SmartDay 的提醒不只弹在电脑上。它可以把同一条提醒同时发到多个地方，任选组合：")
table(["通知方式", "说明", "适合"], [
    ["浏览器通知", "桌面右下角弹窗 + 提示音（默认开启）", "坐在电脑前"],
    ["QQ 机器人（自建 OneBot）", "填服务地址 + QQ 号，直接发私聊消息到你的 QQ", "手机上随时收"],
    ["QQ 官方机器人", "QQ 开放平台机器人，AppID / AppSecret + openid", "已有官方机器人"],
    ["邮箱通知（SMTP）", "填发件邮箱与授权码，邮件到达手机 QQ/微信会提醒", "最省事、最稳定"],
    ["Server 酱 / PushPlus", "填一个 Key，消息推到微信", "不想折腾"],
    ["企业微信 / 钉钉 / 飞书", "群机器人 Webhook，粘贴地址即可", "团队或自用"],
    ["自定义 Webhook", "自填地址与请求体模板，可对接自建服务", "高级用户"],
])
bullets([
    "多选同时收：可同时启用「QQ 机器人 + 邮箱」，一条提醒两边都到；每条通道可单独启用/停用、单独测试；"
    "免打扰时段：默认 23:00–07:00 不推送，可勾选「免打扰时段也推送」；"
    "推送记录：最近 30 条结果留痕，没推成功会写明原因（未开启 / 免打扰 / 平台报错）；"
    "推送内容：事件提醒、任务到期与过期、纪念日与倒计时、专注结束。"
])
figure(os.path.join(SHOTS, "07-push.png"), "图 10  通知设置：多通道可同时启用，每条单独测试，并有推送记录")
h("3.7 设置、搜索与数据管理", 2)
bullets([
    "外观：浅色 / 深色 / 跟随系统；"
    "日历：周起始日、农历/节日/周数开关、工作时段、默认事件时长与默认提醒；"
    "提醒：默认提前分钟、提示音、免打扰时段、通知保留天数；"
    "任务与笔记：新任务默认优先级与位置、完成动效、笔记标签管理；"
    "专注：番茄时长、短休/长休、长休间隔、自动免打扰；"
    "数据管理：一键备份（导出 JSON）、导入恢复、回收站、清空数据；"
    "全局搜索：日历事件、任务、笔记、日记一网打尽；"
    "通知中心：所有提醒集中查看，可标记已读、清除已读。"
])
figure(os.path.join(SHOTS, "06-settings.png"), "图 11  设置页：左侧分组，右侧对应配置项")
doc.add_page_break()

h("四、怎么开始用")
h("方式一：桌面端（推荐，含桌面壁纸日历）", 3)
bullets([
    "双击项目根目录的 start-desktop.bat（或执行 npm start）；"
    "首次启动自动创建数据文件，无需注册、无需联网；"
    "主窗口与壁纸日历共用同一份数据，改动实时同步。"
])
h("方式二：网页版（浏览器直接打开）", 3)
bullets([
    "双击 start-web.bat，或进入 smartday-web 目录执行 npm run dev；"
    "浏览器访问提示的地址即可（默认 http://localhost:5173）；"
    "注意：邮箱通知只能在桌面端发送（浏览器无法直连邮件服务器）。"
])
h("数据与隐私", 2)
bullets([
    "本地优先：所有数据保存在你自己设备的浏览器数据库（IndexedDB）中；"
    "没有服务端：SmartDay 不会把你的日程、任务、笔记上传到任何服务器；"
    "推送是可选功能：只有你主动开启并配置后，提醒内容才会经由你选择的通道发送；"
    "可随时导出：设置 → 数据管理 → 备份，得到完整 JSON，换设备可导入恢复。"
])
h("常见问题", 2)
table(["问题", "解答"], [
    ["需要注册账号吗？", "不需要。安装即用，数据全在本地。"],
    ["手机能用吗？", "暂无手机端。但提醒可以推送到手机（QQ / 邮箱 / 微信）。"],
    ["会不会打扰我？", "可设置免打扰时段，也只有你打开开关后才会推送。"],
    ["桌面日历挡住图标怎么办？", "空白区域点击穿透；也可锁定置底，或调低透明度。"],
    ["换电脑数据怎么办？", "用「数据管理 → 备份」导出 JSON，在新设备导入即可。"],
    ["不写日记会怎样？", "不影响任何其它功能，日记完全可选。"],
])
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run("—— 让每一天都清清楚楚 ——")
r.font.size = Pt(10)
r.font.color.rgb = GREY
doc.save(OUT)
print("SAVED:", OUT)
print("size(KB):", round(os.path.getsize(OUT) / 1024))
